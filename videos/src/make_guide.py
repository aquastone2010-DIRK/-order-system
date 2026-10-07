# -*- coding: utf-8 -*-
"""產生《品富餐飲 店內 AI 短影片｜小白操作手冊》Word 檔。
用法：python make_guide.py KIT_DIR OUT_DOCX   （KIT_DIR = build.py 輸出的 gemini_kit 資料夾）
"""
import glob, json, os, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content import STYLE, VIDEOS  # noqa: E402

FONT = "Microsoft JhengHei"
NAVY, TEAL, RED, GREY = RGBColor(0x1B, 0x36, 0x5D), RGBColor(0x0F, 0x6E, 0x6E), RGBColor(0xB4, 0x23, 0x18), RGBColor(0x66, 0x66, 0x66)


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color is not None: run.font.color.rgb = color


def para(doc, text="", size=10.5, bold=False, color=None, align=None, after=4):
    p = doc.add_paragraph()
    if text:
        set_font(p.add_run(text), size, bold, color)
    p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    return p


def heading(doc, text, level):
    h = doc.add_heading(level=level)
    set_font(h.add_run(text), 16 if level == 1 else 13, True, NAVY if level == 1 else TEAL)
    return h


def bullet(doc, text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    if isinstance(text, (list, tuple)):
        for t in text:
            if isinstance(t, tuple):
                set_font(p.add_run(t[0]), 10.5, True)
            else:
                set_font(p.add_run(t), 10.5)
    else:
        set_font(p.add_run(text), 10.5)
    p.paragraph_format.space_after = Pt(2)
    return p


def shade(cell, hexfill):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexfill)
    tcPr.append(sh)


def table(doc, header, rows, widths, size=9.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        set_font(c.paragraphs[0].add_run(h), size, True, RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, "1B365D")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            set_font(cells[i].paragraphs[0].add_run(str(v)), size)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def box(doc, title, lines, fill="EAF4F4", color=TEAL):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, fill)
    c.text = ""
    set_font(c.paragraphs[0].add_run(title), 11, True, color)
    for l in lines:
        set_font(c.add_paragraph().add_run(l), 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def mmss(t):
    return f"{int(t // 60)}:{t % 60:04.1f}"


def main(kit, out):
    kits = {}
    for p in glob.glob(os.path.join(kit, "*.json")):
        d = json.load(open(p, encoding="utf-8"))
        kits[d["id"]] = d
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10.5)

    # 封面
    para(doc, "", after=80)
    para(doc, "品富餐飲 PINFU", 26, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 6)
    para(doc, "店內 AI 短影片｜小白操作手冊", 22, True, TEAL, WD_ALIGN_PARAGRAPH.CENTER, 20)
    para(doc, "6 部 × 90 秒｜分店電視播放｜用 Gemini 升級成實拍風格", 13, False, None, WD_ALIGN_PARAGRAPH.CENTER, 40)
    table(doc, ["項目", "內容"], [
        ["版本", "V1.0"],
        ["日期", "2026 年 10 月 7 日"],
        ["已完成", "6 部 90 秒成品影片＋9 分鐘循環合輯＋Gemini 製作素材包"],
        ["適用", "童綜合醫院櫃位、科技廠櫃位等所有品富分店"],
    ], [4, 12])
    doc.add_page_break()

    # 1
    heading(doc, "一、你已經拿到的東西（打開就能用）", 1)
    para(doc, "以下影片已經做好，可以直接複製到 USB 隨身碟，插到分店電視播放，不需要任何操作 AI 的技術。")
    rows = [["00", "PINFU_00_店內循環合輯_6部.mp4", "6 部串接", "9:00", "建議分店直接循環播放這一個檔"]]
    for v in VIDEOS:
        rows.append([v["id"], f"PINFU_{v['id']}_{v['title'].replace(' ', '')}.mp4", f"{v['series']}：{v['title']}", "1:30", v["audience"]])
    table(doc, ["#", "檔名", "主題", "長度", "播放對象"], rows, [1, 6.2, 4.3, 1.3, 4.2], 9)
    para(doc, "規格：1920×1080（Full HD）、H.264 影像＋AAC 聲音、MP4。市面上大多數有 USB 播放功能的電視都能直接播放。", 9.5, color=GREY)
    box(doc, "重要說明：這 6 部是「動態字卡版」", [
        "這次交付的影片，畫面由程式自動繪製（圖示＋大字＋字幕），配音使用台灣華語 AI 語音，背景音樂由程式自行合成，沒有版權問題。",
        "我無法登入你的 Google 帳號操作 Gemini，所以「實拍風格」的 AI 畫面需要你依第四章自己產生。",
        "好處是：配音、字幕、分鏡時間我都已經做好放在素材包，你只要把 Gemini 產生的畫面放上去即可。",
    ], "FFF6E5", RGBColor(0x9A, 0x5B, 0x00))

    # 2
    heading(doc, "二、3 個步驟：在分店電視播放", 1)
    heading(doc, "步驟 1：準備 USB 隨身碟", 2)
    bullet(doc, "容量 8GB 以上即可（9 分鐘合輯約 47MB）。")
    bullet(doc, "格式建議 exFAT 或 FAT32（電腦上對隨身碟按右鍵 →「格式化」可以看到）。注意：格式化會清空隨身碟。")
    heading(doc, "步驟 2：複製影片", 2)
    bullet(doc, "把「PINFU_00_店內循環合輯_6部.mp4」複製到隨身碟最外層（不要放太多層資料夾，部分電視找不到）。")
    bullet(doc, "也可以只放單部影片，例如醫院櫃位只放 01、04、06。")
    heading(doc, "步驟 3：電視播放並設定循環", 2)
    bullet(doc, "把隨身碟插進電視背面或側面的 USB 孔。")
    bullet(doc, "按遙控器「輸入／訊號源（Input / Source）」→ 選「USB」或「媒體播放」。")
    bullet(doc, "選影片 → 播放 → 按遙控器「選項／設定（Options / Tools）」→ 找「重複播放」→ 選「全部重複」或「單曲重複」。")
    bullet(doc, "各品牌電視選單名稱不同；找不到時，可上網搜尋「電視型號＋USB 重複播放」。")
    box(doc, "現場小技巧", [
        "• 尖峰時段人聲吵雜：可以靜音播放，所有內容都有大字幕。",
        "• 有聲播放：音量建議先調到能聽清楚、但不干擾點餐對話的大小。",
        "• 電視無法播放：改用電視盒（Android TV Box、Chromecast 等）或筆電接 HDMI 播放。",
    ])

    # 3
    heading(doc, "三、播放排程建議", 1)
    table(doc, ["時段", "建議影片", "原因"], [
        ["開店～早餐", "01 品牌、04 我的餐盤", "建立品牌印象、健康訴求適合早晨"],
        ["午餐尖峰（排隊）", "02 智慧品保標籤、03 食材的旅程", "排隊時最容易看完，建立食安信任"],
        ["下午／晚餐", "05 惜食減廢、06 外帶安心", "外帶比例高，提醒保存方法"],
        ["不想管排程", "00 循環合輯", "9 分鐘一輪，全天循環"],
    ], [3.5, 5.5, 8])

    # 4
    heading(doc, "四、用 Gemini 做「實拍風格」AI 影片（小白版）", 1)
    box(doc, "先知道 3 件事", [
        "1. Gemini 的影片功能（Veo 模型）一次只產生約 8 秒的短片，90 秒影片要分 10 段做，再剪接起來。",
        "2. 產生影片通常需要 Google AI 付費方案，且每日次數有上限；功能名稱、方案與額度以 Google 當下官方說明為準。",
        "3. AI 畫面裡的文字常會變成亂碼，所以提示詞都寫了「no on-screen text」，字幕改用我們做好的 SRT。",
    ], "FDF1F0", RED)
    heading(doc, "4.1 產生每一段畫面", 2)
    steps = [
        "用電腦打開 gemini.google.com，登入你的 Google 帳號。",
        "在對話框下方找到「影片（Video）」按鈕，或從「工具」選單選擇影片／Veo。",
        "打開本手冊第六章，複製某一幕的「Gemini 提示詞」整段（英文），貼進對話框。",
        "按送出，等待約 1～3 分鐘。",
        "預覽：滿意就按下載；不滿意就按重新產生，或在提示詞後面加一句修改要求（例如 make it brighter）。",
        "檔名改成「影片編號-幕次」，例如 02-05.mp4，存在同一個資料夾。",
        "重複以上步驟，直到一部影片的 10 幕全部完成。",
    ]
    for s in steps:
        bullet(doc, s, "List Number")
    heading(doc, "4.2 拼成 90 秒成品（用免費的 CapCut／剪映）", 2)
    steps2 = [
        "下載安裝 CapCut（電腦版），按「新專案」。",
        "匯入 10 段影片，依幕次順序拖到時間軸。",
        "對照第六章表格的「開始～結束」秒數，把每段裁切或拉長到對應長度（片段太短時：選片段 →「變速」調到 0.9 倍）。",
        "匯入素材包的「…_配音音軌.m4a」，拖到音訊軌，從 0 秒開始對齊；把 AI 影片本身的聲音靜音。",
        "選「字幕」→「匯入字幕」→ 選素材包的「….srt」，字幕就會自動對好時間；字體選大一點、加黑色底框。",
        "左上角加上品富 Logo（選用）。",
        "按「匯出」：解析度 1080p、格式 MP4、影格率 25 或 30。",
    ]
    for s in steps2:
        bullet(doc, s, "List Number")
    para(doc, "也可以用 Google Vids 或 Gemini 的 Flow 工具拼接，步驟概念相同：排片段 → 放配音 → 放字幕 → 匯出。", 9.5, color=GREY)

    # 5
    heading(doc, "五、建立你的專屬 Gem「品富短影片導演」", 1)
    para(doc, "Gem 是 Gemini 裡的「客製化助理」。設定一次後，以後只要說「幫我做一部中秋節影片」，它就會用同樣的格式產出腳本與提示詞。")
    for s in ["在 Gemini 左側選單點「Gem 管理員」→「新增 Gem」。",
              "名稱輸入：品富短影片導演。",
              "在「指示」欄貼上下方整段文字 → 儲存。",
              "之後點選這個 Gem 開始對話即可。"]:
        bullet(doc, s, "List Number")
    gem = [
        "你是品富餐飲（PINFU）的店內短影片導演。品富在醫院、科技園區等場域設櫃，提供團膳與餐飲服務。",
        "每次我給你一個主題，請產出一部 90 秒、10 幕的店內播放影片企劃，格式如下：",
        "1. 影片標題、目標觀眾、核心訊息（一句話）。",
        "2. 分鏡表：幕次、秒數（每幕 8～10 秒，總長 90 秒）、畫面大標題（12 字內）、重點字卡（最多 3 行）、台灣華語旁白（每幕 35 字內，口語、親切）。",
        "3. 每幕一段英文 Veo 影片提示詞：16:9、實拍電影感、柔和自然光、乾淨的台灣餐飲廚房或櫃位，畫面中不得出現文字、Logo、品牌名稱。",
        "4. 數據規則：凡是溫度、時間、份量等數字，必須附上來源（例如衛福部國健署、食藥署、食品良好衛生規範準則），沒有可靠來源就不要寫數字。",
        "5. 不得宣稱療效、不得誇大、不得出現其他品牌商標或真實人物。",
        "6. 最後一幕固定為品牌收尾：「品富餐飲」＋一句標語。",
    ]
    box(doc, "Gem 指示（整段複製）", gem, "F6F8FA", NAVY)

    # 6
    heading(doc, "六、6 部影片分鏡表＋Gemini 提示詞", 1)
    para(doc, "每段提示詞 = 該幕畫面描述 + 共用風格句。共用風格句如下，已經自動接在每幕提示詞後面：", 10)
    box(doc, "共用風格句", [STYLE], "F6F8FA", NAVY)
    for v in VIDEOS:
        k = kits[v["id"]]
        doc.add_page_break()
        heading(doc, f"影片 {v['id']}｜{v['title']}（{v['series']}）", 2)
        para(doc, f"播放對象：{v['audience']}　｜　配音：{'女聲' if 'HsiaoChen' in v['voice'] else '男聲'}　｜　素材包檔名開頭：PINFU_{v['id']}_", 9.5, color=GREY)
        rows = []
        for sc in k["scenes"]:
            rows.append([str(sc["scene"]), f"{mmss(sc['start'])}～{mmss(sc['end'])}", sc["head"], sc["say"],
                         sc["veo"] + " " + STYLE])
        table(doc, ["幕", "時間", "畫面標題", "旁白／字幕", "Gemini 提示詞（整段複製）"], rows, [0.8, 2.2, 2.8, 4.6, 6.6], 8)

    # 7
    heading(doc, "七、播放前檢核清單（店長簽核）", 1)
    for s in [
        "影片 01 提到「醫院、科技園區設櫃」與收貨、分區、標示等流程：確認與該分店實際作業一致。",
        "影片 01、02、05 提到智慧品保標籤「正在導入」：若分店尚未導入，可先只播 03、04、06。",
        "AI 實拍版畫面：確認沒有出現其他品牌商標、真實可辨識人物；建議在片尾或角落標示「部分畫面由 AI 生成」。",
        "音量、亮度、播放循環設定完成，並由店長實際看完一輪。",
    ]:
        bullet(doc, s)
    heading(doc, "影片中數據的來源", 2)
    table(doc, ["數據", "出現在", "來源"], [
        ["冷藏 7°C 以下、冷凍 −18°C 以下、熱藏 60°C 以上", "03、06", "衛生福利部「食品良好衛生規範準則」"],
        ["我的餐盤六口訣、乳品每日 1.5～2 杯（1 杯 240 毫升）、深色蔬菜 1/3、全穀 1/3 未精製、豆魚蛋肉優先順序、堅果種子每日 1 份", "04", "衛生福利部國民健康署「我的餐盤」與「每日飲食指南」"],
        ["熟食室溫不超過 2 小時；高於 32°C 不超過 1 小時；復熱中心溫度 74°C；分小份淺盒降溫", "06", "美國農業部食品安全檢驗局（USDA FSIS）剩食處理指引；食藥署食品保存衛教"],
        ["櫃位使用期限不得晚於原包裝有效日", "02", "品富 FoodOps 內部 SOP（見專案計畫書）"],
    ], [7, 1.6, 8.4], 9)
    para(doc, "若主管機關日後更新數據，請修改 videos/src/content.py 後重新執行產生程式，影片會自動重製。", 9.5, color=GREY)

    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
